"""Extração de texto com referências de página/slide. Extratores desacoplados por tipo.

``sniff_kind`` valida o tipo real pelos bytes (uma página HTML de login servida como
"PDF" é detectada como html e rejeitada como material).
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser

EXTRACTOR_VERSION = "1"

MIN_CHARS_PER_PAGE = 20  # abaixo disso a página conta como "sem texto"


@dataclass
class Segment:
    ref: str  # "p.3", "slide 5", "seção: Título"
    text: str


@dataclass
class Extraction:
    status: str  # extracted | no_text | unsupported | failed
    kind: str
    segments: list[Segment] = field(default_factory=list)
    page_count: int = 0
    warnings: list[str] = field(default_factory=list)
    error: str | None = None

    @property
    def char_count(self) -> int:
        return sum(len(s.text) for s in self.segments)


def sniff_kind(data: bytes, declared_name: str | None = None, content_type: str | None = None) -> str:
    head = data[:2048].lstrip()
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = z.namelist()
            if any(n.startswith("ppt/") for n in names):
                return "pptx"
            if any(n.startswith("word/") for n in names):
                return "docx"
        except zipfile.BadZipFile:
            pass
        return "zip"
    low = head[:512].lower()
    if low.startswith(b"<!doctype html") or b"<html" in low or b"<head" in low or b"<body" in low:
        return "html"
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        return "ole"  # .ppt/.doc antigos
    ext = (declared_name or "").rsplit(".", 1)[-1].lower() if declared_name and "." in declared_name else ""
    try:
        data[:4096].decode("utf-8")
        is_text = True
    except UnicodeDecodeError:
        is_text = False
    if is_text:
        if ext in ("md", "markdown"):
            return "md"
        if ext in ("txt", "csv", "log") or (content_type or "").startswith("text/"):
            return "txt"
        return "md" if ext == "md" else "txt"
    return "binary"


def is_login_page(data: bytes) -> bool:
    low = data[:8192].lower()
    if b"<form" not in low:
        return False
    return any(
        k in low
        for k in (b"type='password'", b'type="password"', b"logar.do", b"entrar no sistema", b"login")
    )


def extract(data: bytes, kind: str) -> Extraction:
    try:
        if kind == "pdf":
            return _extract_pdf(data)
        if kind == "pptx":
            return _extract_pptx(data)
        if kind == "html":
            return _extract_html(data)
        if kind in ("txt", "md"):
            return _extract_text(data, kind)
        reason = {
            "docx": "DOCX ainda não suportado",
            "ole": "formato antigo (.ppt/.doc) exige conversão",
            "zip": "arquivo compactado",
            "binary": "binário desconhecido",
        }.get(kind, f"tipo não suportado: {kind}")
        return Extraction(status="unsupported", kind=kind, error=reason)
    except Exception as e:  # extrator quebrado não derruba o pipeline
        return Extraction(status="failed", kind=kind, error=f"{type(e).__name__}: {e}"[:300])


def _extract_pdf(data: bytes) -> Extraction:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            return Extraction(status="failed", kind="pdf", error="PDF criptografado")
    segs: list[Segment] = []
    empty_pages = 0
    for i, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        if len(text) < MIN_CHARS_PER_PAGE:
            empty_pages += 1
            continue
        segs.append(Segment(ref=f"p.{i}", text=text))
    n = len(reader.pages)
    warnings = []
    if empty_pages:
        warnings.append(f"{empty_pages}/{n} página(s) sem camada de texto (escaneado/imagem?)")
    if not segs:
        return Extraction(
            status="no_text",
            kind="pdf",
            page_count=n,
            warnings=warnings,
            error="PDF sem camada de texto; exige OCR/conversão (não executado)",
        )
    return Extraction(status="extracted", kind="pdf", segments=segs, page_count=n, warnings=warnings)


def _extract_pptx(data: bytes) -> Extraction:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    segs: list[Segment] = []
    thin = 0
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[notas] {notes}")
        text = _clean("\n".join(parts))
        if len(text) < MIN_CHARS_PER_PAGE:
            thin += 1
            if text:
                segs.append(Segment(ref=f"slide {i}", text=text + " [slide com pouco texto; ver original]"))
            continue
        segs.append(Segment(ref=f"slide {i}", text=text))
    n = len(prs.slides)
    warnings = [f"{thin}/{n} slide(s) com pouco texto (figuras/diagramas?)"] if thin else []
    if not any(len(s.text) >= MIN_CHARS_PER_PAGE for s in segs):
        return Extraction(
            status="no_text", kind="pptx", page_count=n, warnings=warnings, error="PPTX sem texto útil"
        )
    return Extraction(status="extracted", kind="pptx", segments=segs, page_count=n, warnings=warnings)


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript", "head", "title"}
    HEAD = {"h1", "h2", "h3", "h4"}

    def __init__(self):
        super().__init__()
        self.sections: list[tuple[str, list[str]]] = [("início", [])]
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag in self.HEAD:
            self.sections.append(("", []))
            self._in_head = True
        elif tag in ("p", "li", "br", "div", "tr"):
            self.sections[-1][1].append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        elif tag in self.HEAD:
            self._in_head = False

    def handle_data(self, data):
        if self._skip:
            return
        if getattr(self, "_in_head", False) and not self.sections[-1][0]:
            title, body = self.sections[-1]
            self.sections[-1] = ((title + data).strip(), body)
        else:
            self.sections[-1][1].append(data)


def _extract_html(data: bytes) -> Extraction:
    text = data.decode("utf-8", errors="replace")
    p = _HTMLText()
    p.feed(text)
    segs = []
    for title, parts in p.sections:
        body = _clean("".join(parts))
        if not body and not title:
            continue
        segs.append(Segment(ref=f"seção: {title or 'início'}", text=(title + "\n" + body).strip()))
    if not segs or sum(len(s.text) for s in segs) < MIN_CHARS_PER_PAGE:
        return Extraction(status="no_text", kind="html", error="HTML sem texto útil")
    return Extraction(status="extracted", kind="html", segments=segs, page_count=len(segs))


def _extract_text(data: bytes, kind: str) -> Extraction:
    text = data.decode("utf-8", errors="replace")
    if kind == "md":
        segs = []
        cur_title, buf = "início", []
        for line in text.splitlines():
            m = re.match(r"^(#{1,4})\s+(.*)", line)
            if m:
                if buf and "".join(buf).strip():
                    segs.append(Segment(ref=f"seção: {cur_title}", text=_clean("\n".join(buf))))
                cur_title, buf = m.group(2).strip(), [line]
            else:
                buf.append(line)
        if buf and "".join(buf).strip():
            segs.append(Segment(ref=f"seção: {cur_title}", text=_clean("\n".join(buf))))
    else:
        segs = [Segment(ref="texto", text=_clean(text))]
    if not segs or sum(len(s.text) for s in segs) < MIN_CHARS_PER_PAGE:
        return Extraction(status="no_text", kind=kind, error="arquivo de texto vazio")
    return Extraction(status="extracted", kind=kind, segments=segs, page_count=len(segs))


def _clean(text: str) -> str:
    text = text.replace("\r", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
